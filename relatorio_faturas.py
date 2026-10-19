#!/usr/bin/env python3
"""
Relatório de Faturas do Atacado - TOTAL E BAIXAS
=================================================

Pedido de Laura e Janaína:
  1. Número da transação MESMO APÓS A BAIXA (parcial ou total)
  2. Valor vendido MESMO COM BAIXAS (o valor original da fatura, não só o saldo)

Problema no TOTVS: na baixa parcial a fatura é dividida (parcela 1 = parte paga,
parcela 200+ = saldo) e o saldo fica com a transação zerada no relatório 0433.
Aqui cada fatura é remontada a partir de todas as suas parcelas:

  Valor vendido   = soma de todas as parcelas da fatura (original)
  Total baixado   = soma das parcelas já baixadas
  Saldo em aberto = soma das parcelas ainda não baixadas
  Transação       = transação da NF de venda (vale para todas as parcelas)

Uso:
  python relatorio_faturas.py                         # vencimento do .env (VENC_INI/VENC_FIM)
  python relatorio_faturas.py --venc-ini 2026-09-01 --venc-fim 2026-11-30

Saída: saida_api/faturas_atacado_<data>.xlsx com as abas
  Em aberto | Todas as faturas | Baixas (detalhe) | Créditos devolução | Resumo por corretor
"""

from __future__ import annotations

import argparse
import calendar
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Instale o openpyxl:  pip install openpyxl")

ap = argparse.ArgumentParser(description="Relatório de faturas do atacado (total e baixas)")
ap.add_argument("--venc-ini", help="vencimento inicial AAAA-MM-DD (padrão: VENC_INI do .env)")
ap.add_argument("--venc-fim", help="vencimento final AAAA-MM-DD (padrão: VENC_FIM do .env)")
ARGS, _ = ap.parse_known_args()  # parse_known_args: permite ser importado por outros scripts
if ARGS.venc_ini:
    os.environ["VENC_INI"] = ARGS.venc_ini
if ARGS.venc_fim:
    os.environ["VENC_FIM"] = ARGS.venc_fim

import sonda_api_totvs as S  # noqa: E402
import reconstruir_0433 as R  # noqa: E402


# ----------------------------------------------------------------------------
# Coleta (mesma lógica validada no reconstruir_0433: bate no centavo com a planilha)
# ----------------------------------------------------------------------------

def limite_6_meses(ini: date) -> date:
    """Última data aceita pela API para um período que começa em `ini` (6 meses - 1 dia)."""
    y, m = divmod(ini.month - 1 + 6, 12)
    y, m = ini.year + y, m + 1
    return date(y, m, min(ini.day, calendar.monthrange(y, m)[1])) - timedelta(days=1)


def janelas_6_meses(ini: date, fim: date) -> list[tuple[date, date]]:
    janelas = []
    while ini <= fim:
        ate = min(limite_6_meses(ini), fim)
        janelas.append((ini, ate))
        ini = ate + timedelta(days=1)
    return janelas


def coletar() -> list[dict]:
    api = S.Api()
    api.autenticar()
    print(f"\n▶ Títulos com vencimento de {S.VENC_INI} a {S.VENC_FIM}")
    titulos = R.buscar_tudo(api, R.PATH_AR, {
        "branchCodeList": S.BRANCHES,
        "startExpiredDate": S.data_iso(S.VENC_INI),
        "endExpiredDate": S.data_iso(S.VENC_FIM, fim=True),
    }, "invoice,calculateValue", "títulos")
    titulos = [t for t in titulos if t.get("status") not in R.STATUS_EXCLUIR]

    emissoes = sorted(t["issueDate"][:10] for t in titulos if t.get("issueDate"))
    if not emissoes:
        sys.exit("Nenhum título no período.")
    print(f"\n▶ Notas de atacado emitidas de {emissoes[0]} a {emissoes[-1]}")
    janelas = janelas_6_meses(date.fromisoformat(emissoes[0]), date.fromisoformat(emissoes[-1]))
    notas = []
    for i, (ini, fim) in enumerate(janelas, 1):  # a API só aceita até 6 meses de emissão por consulta
        notas += R.buscar_tudo(api, R.PATH_NF, {
            "branchCodeList": S.BRANCHES,
            "operationCodeList": sorted(R.OPERACOES_ATACADO),
            "startIssueDate": S.data_iso(ini.isoformat()),
            "endIssueDate": S.data_iso(fim.isoformat(), fim=True),
        }, None, "notas" if len(janelas) == 1 else f"notas {i}/{len(janelas)}")
    notas = [n for n in notas if n.get("invoiceStatus") not in ("Canceled", "Deleted", "Denied")]

    clientes = {n.get("personCode") for n in notas}
    print(f"\n▶ Cadastro de {len(clientes)} clientes")
    pessoas = R.buscar_pessoas(api, clientes)
    linhas, _ = R.montar(titulos, notas, pessoas)
    return linhas


# ----------------------------------------------------------------------------
# Agrupamento por fatura
# ----------------------------------------------------------------------------

def br_para_data(s: str):
    if not s:
        return None
    try:
        d, m, y = s.split("/")
        return date(int(y), int(m), int(d))
    except ValueError:
        return s


def agrupar(linhas: list[dict]):
    faturas: dict[tuple, dict] = {}
    baixas: list[dict] = []
    for r in linhas:
        chave = (r["CD cliente"], r["Nr. Fat"], r["Natureza"])
        f = faturas.setdefault(chave, {
            "Corretor": r["Corretor / Representante"],
            "Empresa": r["Empresa"],
            "Cliente": r["Cliente"],
            "CD cliente": r["CD cliente"],
            "CPF/CNPJ": r["Cpfcnpj"],
            "Nr. Fat / NF": r["Nr. Fat"],
            "Transação": r["Transação"],
            "Natureza": r["Natureza"],
            "Emissão": br_para_data(r["Emissao"]),
            "Vencimento": br_para_data(r["Vencimento"]),
            "Prz": int(r["Prz"]) if str(r["Prz"]).lstrip("-").isdigit() else r["Prz"],
            "Cód. Oper": r.get("Cod. Oper"),
            "Status cliente": r.get("Status cliente"),
            "Valor vendido": 0.0,
            "Total baixado": 0.0,
            "Valor recebido": 0.0,
            "Saldo em aberto": 0.0,
            "Qtd parcelas": 0,
            "Qtd baixas": 0,
            "Última baixa": None,
            "Valor total NF": r["Valor total NF"],
        })
        v = float(r["Valor"] or 0)
        f["Valor vendido"] += v
        f["Qtd parcelas"] += 1
        if str(r["Tp. Baixa"]) == "0":
            f["Saldo em aberto"] += v
        else:
            f["Total baixado"] += v
            f["Valor recebido"] += float(r["Valor pago"] or 0)
            f["Qtd baixas"] += 1
            dt = br_para_data(r["Dt. Baixa"])
            if dt and (f["Última baixa"] is None or dt > f["Última baixa"]):
                f["Última baixa"] = dt
            baixas.append({
                "Corretor": r["Corretor / Representante"], "Cliente": r["Cliente"],
                "CD cliente": r["CD cliente"], "Nr. Fat / NF": r["Nr. Fat"],
                "Transação": r["Transação"], "Parcela": r["Parcela"],
                "Parcela origem": r["Parcela origem"] or None,
                "Data baixa": dt, "Valor baixado": v, "Valor recebido": float(r["Valor pago"] or 0),
                "Tipo de baixa": r["Tipo baixa (nome)"], "Natureza": r["Natureza"],
            })
    for f in faturas.values():
        for k in ("Valor vendido", "Total baixado", "Valor recebido", "Saldo em aberto"):
            f[k] = round(f[k], 2)
        if f["Saldo em aberto"] <= 0.005:
            f["Situação"] = "Quitada"
        elif f["Total baixado"] > 0:
            f["Situação"] = "Parcialmente paga"
        else:
            f["Situação"] = "Em aberto"
    ordem = lambda f: (f["Corretor"], str(f["Cliente"]), str(f["Nr. Fat / NF"]))  # noqa: E731
    return sorted(faturas.values(), key=ordem), sorted(baixas, key=lambda b: (b["Corretor"], str(b["Cliente"]),
                                                                               str(b["Nr. Fat / NF"]), b["Parcela"] or 0))


# ----------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------

COLS_FATURA = ["Corretor", "Cliente", "CD cliente", "CPF/CNPJ", "Status cliente", "Nr. Fat / NF", "Transação",
               "Emissão", "Vencimento", "Prz", "Cód. Oper", "Valor vendido", "Total baixado", "Saldo em aberto",
               "Situação", "Qtd baixas", "Última baixa", "Valor recebido", "Qtd parcelas"]
COLS_BAIXA = ["Corretor", "Cliente", "CD cliente", "Nr. Fat / NF", "Transação", "Parcela",
              "Parcela origem", "Data baixa", "Valor baixado", "Valor recebido", "Tipo de baixa"]
MOEDA = '"R$" #,##0.00'
DATA = "DD/MM/YYYY"
CAB = PatternFill("solid", fgColor="1F4E78")
FAIXA_PARCIAL = PatternFill("solid", fgColor="FFF2CC")


def escrever_aba(wb, titulo: str, colunas: list[str], linhas: list[dict], destaque_parcial=False,
                 total_cols=()):
    ws = wb.create_sheet(titulo)
    ws.append(colunas)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = CAB
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in linhas:
        ws.append([r.get(c) for c in colunas])
        if destaque_parcial and r.get("Situação") == "Parcialmente paga":
            for c in ws[ws.max_row]:
                c.fill = FAIXA_PARCIAL
    n = ws.max_row
    if linhas and total_cols:
        ws.append([])
        ws.append(["TOTAL"] + [None] * (len(colunas) - 1))
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for nome in total_cols:
            j = colunas.index(nome) + 1
            L = get_column_letter(j)
            cel = ws.cell(ws.max_row, j, f"=SUBTOTAL(9,{L}2:{L}{n})")
            cel.font = Font(bold=True)
            cel.number_format = MOEDA
    for j, nome in enumerate(colunas, 1):
        L = get_column_letter(j)
        fmt = MOEDA if nome.startswith(("Valor", "Total", "Saldo")) else DATA if nome in (
            "Emissão", "Vencimento", "Última baixa", "Data baixa") else None
        if fmt:
            for i in range(2, n + 1):
                ws[f"{L}{i}"].number_format = fmt
        larg = max([len(str(nome))] + [len(str(r.get(nome) or "")) for r in linhas[:500]])
        ws.column_dimensions[L].width = min(max(10, larg + 2), 42)
    ws.freeze_panes = "C2"
    if linhas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{n}"
    return ws


def gerar_excel(faturas: list[dict], baixas: list[dict]) -> Path:
    vendas = [f for f in faturas if f["Natureza"] == "Venda"]
    creditos = [f for f in faturas if f["Natureza"] != "Venda"]
    abertas = [f for f in vendas if f["Situação"] != "Quitada"]

    wb = Workbook()
    capa = wb.active
    capa.title = "Leia-me"
    info = [
        ["Relatório de Faturas do Atacado - TOTAL E BAIXAS"],
        [f"Gerado em {datetime.now():%d/%m/%Y %H:%M} direto da API TOTVS Moda"],
        [f"Vencimento: {S.VENC_INI} a {S.VENC_FIM} | Empresa {S.BRANCHES} | Operações {sorted(R.OPERACOES_ATACADO)}"],
        [],
        ["Como ler"],
        ["• Transação: número da transação da venda, mantido mesmo depois de baixas parciais ou totais."],
        ["• Valor vendido: valor original da fatura (soma de todas as parcelas, inclusive as já pagas)."],
        ["• Total baixado: quanto da fatura já foi baixado. Saldo em aberto: quanto falta."],
        ["• Linhas em amarelo: faturas parcialmente pagas (antes apareciam com transação 0)."],
        ["• Títulos cancelados no TOTVS não entram. Créditos de devolução (Credev) ficam em aba própria."],
        [],
        ["Resumo", "Qtd", "Valor"],
        ["Faturas de venda", len(vendas), round(sum(f["Valor vendido"] for f in vendas), 2)],
        ["  já baixado", None, round(sum(f["Total baixado"] for f in vendas), 2)],
        ["  em aberto", len(abertas), round(sum(f["Saldo em aberto"] for f in vendas), 2)],
        ["  parcialmente pagas", sum(1 for f in vendas if f["Situação"] == "Parcialmente paga"), None],
        ["Créditos de devolução", len(creditos), round(sum(f["Valor vendido"] for f in creditos), 2)],
    ]
    for linha in info:
        capa.append(linha)
    capa["A1"].font = Font(bold=True, size=14)
    capa["A5"].font = Font(bold=True)
    for c in capa[12]:
        c.font = Font(bold=True)
    for i in range(13, 18):
        capa[f"C{i}"].number_format = MOEDA
    capa.column_dimensions["A"].width = 30
    capa.column_dimensions["C"].width = 18

    totais = ("Valor vendido", "Total baixado", "Saldo em aberto", "Valor recebido")
    escrever_aba(wb, "Em aberto", COLS_FATURA, abertas, destaque_parcial=True, total_cols=totais)
    escrever_aba(wb, "Todas as faturas", COLS_FATURA, vendas, destaque_parcial=True, total_cols=totais)
    escrever_aba(wb, "Baixas (detalhe)", COLS_BAIXA, [b for b in baixas if b["Natureza"] == "Venda"],
                 total_cols=("Valor baixado", "Valor recebido"))
    escrever_aba(wb, "Créditos devolução", COLS_FATURA, creditos, total_cols=totais)

    resumo = defaultdict(lambda: {"Faturas": 0, "Valor vendido": 0.0, "Total baixado": 0.0,
                                  "Saldo em aberto": 0.0, "Faturas em aberto": 0})
    for f in vendas:
        r = resumo[f["Corretor"] or "(sem corretor)"]
        r["Faturas"] += 1
        r["Valor vendido"] += f["Valor vendido"]
        r["Total baixado"] += f["Total baixado"]
        r["Saldo em aberto"] += f["Saldo em aberto"]
        r["Faturas em aberto"] += f["Situação"] != "Quitada"
    linhas_res = [{"Corretor": k, **{kk: (round(vv, 2) if isinstance(vv, float) else vv) for kk, vv in v.items()}}
                  for k, v in sorted(resumo.items())]
    escrever_aba(wb, "Resumo por corretor",
                 ["Corretor", "Faturas", "Faturas em aberto", "Valor vendido", "Total baixado", "Saldo em aberto"],
                 linhas_res, total_cols=("Valor vendido", "Total baixado", "Saldo em aberto"))

    S.SAIDA.mkdir(exist_ok=True)
    arq = S.SAIDA / f"faturas_atacado_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    wb.save(arq)
    return arq


def main() -> None:
    linhas = coletar()
    faturas, baixas = agrupar(linhas)
    arq = gerar_excel(faturas, baixas)
    vendas = [f for f in faturas if f["Natureza"] == "Venda"]
    print(f"\n✔ {len(vendas)} faturas de venda | vendido R$ {sum(f['Valor vendido'] for f in vendas):,.2f} | "
          f"em aberto R$ {sum(f['Saldo em aberto'] for f in vendas):,.2f} | "
          f"{sum(1 for f in vendas if f['Situação'] == 'Parcialmente paga')} parcialmente pagas")
    print(f"  Excel: {arq}")


if __name__ == "__main__":
    main()
